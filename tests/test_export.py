from pathlib import Path

import pytest

from matera.core.contracts import AnswerKey, NormalizedAnswer, NormalizedPageResult, ReviewTask
from matera.core.profile import FormProfile, OptionDef, QuestionDef
from matera.export.excel import export_to_excel, flatten_result, generate_headers


@pytest.fixture
def sample_profile() -> FormProfile:
    return FormProfile(
        form_id="test",
        form_version="v1",
        questions=(
            QuestionDef(
                question_id="q1",
                response_type="single_select",
                mark_strategy="circle",
                options=(OptionDef("o1"), OptionDef("o2")),
                max_selections=1,
            ),
        ),
    )


def test_generate_headers(sample_profile: FormProfile):
    headers = generate_headers(sample_profile)
    assert headers == [
        "form_id",
        "form_version",
        "page_number",
        "page_status",
        "q1_o1",
        "q1_o2",
        "review_tasks",
    ]


def test_flatten_result_resolved(sample_profile: FormProfile):
    ans1 = NormalizedAnswer(
        answer_key=AnswerKey("test", "v1", 1, "q1", "o1"),
        selected=True,
        resolution_status="resolved",
        decision_source="deterministic",
    )
    ans2 = NormalizedAnswer(
        answer_key=AnswerKey("test", "v1", 1, "q1", "o2"),
        selected=False,
        resolution_status="resolved",
        decision_source="deterministic",
    )
    result = NormalizedPageResult(
        form_id="test",
        form_version="v1",
        page_number=1,
        answers=(ans1, ans2),
    )

    row = flatten_result(result, sample_profile)
    assert row == ["test", "v1", 1, "resolved", 1, 0, ""]


def test_flatten_result_needs_review(sample_profile: FormProfile):
    ans1 = NormalizedAnswer(
        answer_key=AnswerKey("test", "v1", 1, "q1", "o1"),
        selected=None,
        resolution_status="needs_review",
        decision_source="deterministic",
    )
    ans2 = NormalizedAnswer(
        answer_key=AnswerKey("test", "v1", 1, "q1", "o2"),
        selected=None,
        resolution_status="needs_review",
        decision_source="deterministic",
    )
    rt1 = ReviewTask(
        task_id="t1",
        answer_key=ans1.answer_key,
        reason="Ambiguous",
        status="pending",
    )
    rt2 = ReviewTask(
        task_id="t2",
        answer_key=ans2.answer_key,
        reason="Ambiguous",
        status="pending",
    )

    result = NormalizedPageResult(
        form_id="test",
        form_version="v1",
        page_number=1,
        answers=(ans1, ans2),
        review_tasks=(rt1, rt2),
    )

    row = flatten_result(result, sample_profile)

    # Needs review -> empty string ""
    assert row[4] == ""
    assert row[5] == ""

    # review_tasks JSON verification
    import json

    tasks = json.loads(row[6])
    assert len(tasks) == 2
    assert tasks[0]["question"] == "q1"
    assert tasks[0]["option"] == "o1"
    assert tasks[0]["reason"] == "Ambiguous"
    assert tasks[1]["question"] == "q1"
    assert tasks[1]["option"] == "o2"
    assert tasks[1]["reason"] == "Ambiguous"


def test_export_to_excel_end_to_end(tmp_path: Path, sample_profile: FormProfile):
    # Setup results
    ans1 = NormalizedAnswer(
        answer_key=AnswerKey("test", "v1", 1, "q1", "o1"),
        selected=True,
        resolution_status="resolved",
        decision_source="deterministic",
    )
    ans2 = NormalizedAnswer(
        answer_key=AnswerKey("test", "v1", 1, "q1", "o2"),
        selected=False,
        resolution_status="resolved",
        decision_source="deterministic",
    )
    result = NormalizedPageResult(
        form_id="test",
        form_version="v1",
        page_number=1,
        answers=(ans1, ans2),
    )

    output_file = tmp_path / "out.xlsx"
    export_to_excel([result], sample_profile, output_file)

    assert output_file.exists()

    # Verify the contents by reading it back
    import openpyxl

    wb = openpyxl.load_workbook(output_file)
    ws = wb.active

    # Check headers (row 1)
    header_row = [cell.value for cell in ws[1]]
    assert header_row == generate_headers(sample_profile)

    # Check data (row 2)
    # openpyxl reads empty cells as None, but our flat list puts ""
    data_row = [cell.value if cell.value is not None else "" for cell in ws[2]]

    expected_row = flatten_result(result, sample_profile)
    assert data_row == expected_row

    # Will fail until implemented, but confirms import works.
